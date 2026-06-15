"""
Jump Shop quick checkout runner.

Reads one cookie string per Excel row, then runs cart-add -> checkout -> address/card
fill -> payment submit directly with Playwright. This script is intentionally
standalone and does not use the backend batch task flow.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from jump_shop import (
    DEFAULT_USER_AGENT,
    JUMP_SHOP_BASE_URL,
    JUMP_SHOP_CHECKOUT_URL,
    browser_executable_path,
    build_expiry_value,
    compact_text,
    detect_jump_shop_3ds,
    detect_jump_shop_success,
    extract_order_number,
    fetch_jump_shop_product_info,
    fill_billing_address,
    fill_card_fields,
    fill_shipping_address,
    launch_browser_context,
    maybe_fill,
    normalize_jump_shop_product_url,
)


RESULT_FIELDS = ["row", "email", "success", "status", "message", "orderNo", "finalUrl"]
COOKIE_FALSE_VALUES = {"false", "0", "no", "n", "off", "disabled", "skip", "否", "跳过"}
PRINT_LOCK = threading.Lock()


@dataclass(frozen=True)
class AccountCookie:
    row: int
    cookie_header: str
    email: str = ""


@dataclass(frozen=True)
class ProductConfig:
    product_url: str
    variant_id: int
    product_id: int
    section_id: str
    quantity: int
    title: str = ""
    expected_title_contains: str = ""


def log(message: str) -> None:
    with PRINT_LOCK:
        print(f"[{datetime.now().strftime('%H:%M:%S')}] {message}", flush=True)


def load_json(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8-sig") as file:
        data = json.load(file)
    if not isinstance(data, dict):
        raise ValueError("config 必须是 JSON object")
    return data


def resolve_path(base_dir: Path, value: Any) -> Path:
    text = str(value or "").strip()
    if not text:
        raise ValueError("路径不能为空")
    path = Path(text)
    return path if path.is_absolute() else base_dir / path


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def string_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_enabled(value: Any) -> bool:
    text = string_value(value).lower()
    return text not in COOKIE_FALSE_VALUES


def load_accounts_from_excel(path: Path, sheet_name: str = "") -> List[AccountCookie]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先执行: pip install -r requirements.txt") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        return []

    first_row = [normalize_header(value) for value in rows[0]]
    has_header = any(name in first_row for name in ["cookie", "cookieheader", "cookies"])
    accounts: List[AccountCookie] = []

    if has_header:
        header = {name: index for index, name in enumerate(first_row) if name}
        cookie_index = next(
            (header[name] for name in ["cookie", "cookieheader", "cookies"] if name in header),
            None,
        )
        if cookie_index is None:
            raise ValueError("Excel 表头存在但缺少 cookie 列")
        email_index = next((header[name] for name in ["email", "mail", "account"] if name in header), None)
        enabled_index = next((header[name] for name in ["enabled", "enable", "active"] if name in header), None)
        data_rows = rows[1:]
        row_offset = 2
    else:
        cookie_index = 0
        email_index = 1 if len(rows[0]) > 1 else None
        enabled_index = None
        data_rows = rows
        row_offset = 1

    for offset, row in enumerate(data_rows, row_offset):
        cookie_header = string_value(row[cookie_index] if cookie_index < len(row) else "")
        if not cookie_header:
            continue
        if enabled_index is not None and enabled_index < len(row) and not is_enabled(row[enabled_index]):
            continue
        email = string_value(row[email_index] if email_index is not None and email_index < len(row) else "")
        accounts.append(AccountCookie(row=offset, cookie_header=cookie_header, email=email))
    return accounts


def parse_cookie_header(cookie_header: str) -> List[Tuple[str, str]]:
    cookies: List[Tuple[str, str]] = []
    seen = set()
    for part in str(cookie_header or "").split(";"):
        if "=" not in part:
            continue
        name, value = part.split("=", 1)
        name = name.strip()
        value = value.strip()
        if not name or name.lower() in {"path", "domain", "expires", "max-age", "samesite", "secure", "httponly"}:
            continue
        if name in seen:
            continue
        seen.add(name)
        cookies.append((name, value))
    return cookies


def add_cookie_header_to_context(context: Any, cookie_header: str) -> int:
    cookie_items = []
    for name, value in parse_cookie_header(cookie_header):
        item: Dict[str, Any] = {
            "name": name,
            "value": value,
            "path": "/",
            "secure": True,
            "sameSite": "Lax",
        }
        if name.startswith("__Host-"):
            item["url"] = JUMP_SHOP_BASE_URL
        else:
            item["domain"] = ".jumpshop-benelic.com"
        cookie_items.append(item)
    if cookie_items:
        context.add_cookies(cookie_items)
    return len(cookie_items)


def require_int(value: Any, name: str) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name} 必须是数字") from exc
    if number <= 0:
        raise ValueError(f"{name} 必须大于 0")
    return number


def normalize_match_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().lower())


def title_matches(actual: Any, expected: Any) -> bool:
    actual_text = normalize_match_text(actual)
    expected_text = normalize_match_text(expected)
    if not expected_text:
        return True
    if not actual_text:
        return False
    return expected_text in actual_text or actual_text in expected_text


def resolve_product(config: Dict[str, Any]) -> ProductConfig:
    product_url = normalize_jump_shop_product_url(str(config.get("productUrl") or config.get("eventUrl") or ""))
    if not product_url:
        raise ValueError("config 缺少 productUrl")

    parsed: Dict[str, Any] = {}
    try:
        log(f"[product] 解析商品页: {product_url}")
        parsed = fetch_jump_shop_product_info(product_url)
        log(
            "[product] 解析成功: "
            f"title={parsed.get('title') or '-'} variantId={parsed.get('variantId')} "
            f"productId={parsed.get('productId')} sectionId={parsed.get('sectionId')}"
        )
    except Exception as exc:
        log(f"[product] 自动解析失败，将尝试使用 config 手动参数: {exc}")

    variant_id = config.get("variantId") or parsed.get("variantId")
    product_id = config.get("productId") or parsed.get("productId")
    section_id = str(config.get("sectionId") or parsed.get("sectionId") or "").strip()
    if not variant_id or not product_id or not section_id:
        raise ValueError("商品缺少 variantId/productId/sectionId，请在 config 中手动补齐")
    parsed_title = str(parsed.get("title") or "").strip()
    configured_expected_title = str(config.get("expectedProductTitleContains") or config.get("expectedTitleContains") or "").strip()
    if bool(config.get("submitPayment", True)) and not configured_expected_title:
        raise ValueError("submitPayment=true 时必须手动设置 expectedProductTitleContains，防止商品 URL 填错后直接付款")
    expected_title = configured_expected_title or parsed_title
    if configured_expected_title and parsed_title and not title_matches(parsed_title, configured_expected_title):
        raise ValueError(f"商品标题与 expectedProductTitleContains 不匹配: title={parsed_title}")

    return ProductConfig(
        product_url=product_url,
        variant_id=require_int(variant_id, "variantId"),
        product_id=require_int(product_id, "productId"),
        section_id=section_id,
        quantity=require_int(config.get("quantity") or 1, "quantity"),
        title=parsed_title,
        expected_title_contains=expected_title,
    )


def validate_profile(profile: Dict[str, Any]) -> None:
    required = [
        "lastName",
        "firstName",
        "phone",
        "postalCode",
        "province",
        "city",
        "address1",
        "cardHolderName",
        "cardNumber",
        "expMonth",
        "expYear",
        "cvv",
    ]
    missing = [key for key in required if not string_value(profile.get(key))]
    if missing:
        raise ValueError("profile 缺少字段: " + ", ".join(missing))
    build_expiry_value(profile.get("expMonth"), profile.get("expYear"))
    if not profile.get("billingSameAsShipping", True):
        billing_required = ["billingLastName", "billingFirstName", "billingPhone", "billingPostalCode", "billingProvince", "billingCity", "billingAddress1"]
        billing_missing = [key for key in billing_required if not string_value(profile.get(key))]
        if billing_missing:
            raise ValueError("profile 缺少账单地址字段: " + ", ".join(billing_missing))


def create_browser(playwright: Any, headless: bool) -> Any:
    launch_options: Dict[str, Any] = {"headless": bool(headless)}
    browser_path = browser_executable_path()
    if browser_path:
        launch_options["executable_path"] = browser_path
    return playwright.chromium.launch(**launch_options)


def add_to_cart(page: Any, product: ProductConfig) -> Dict[str, Any]:
    return page.evaluate(
        """
        async ({ variantId, quantity, productId, sectionId }) => {
          const formData = new FormData();
          formData.append('form_type', 'product');
          formData.append('utf8', '✓');
          formData.append('quantity', String(quantity));
          formData.append('id', String(variantId));
          formData.append('product-id', String(productId));
          formData.append('section-id', String(sectionId));
          const response = await fetch('/cart/add.js', {
            method: 'POST',
            credentials: 'include',
            headers: { 'X-Requested-With': 'XMLHttpRequest' },
            body: formData
          });
          const text = await response.text();
          return { ok: response.ok, status: response.status, text };
        }
        """,
        {
            "variantId": product.variant_id,
            "quantity": product.quantity,
            "productId": product.product_id,
            "sectionId": product.section_id,
        },
    )


def clear_cart(page: Any) -> Dict[str, Any]:
    return page.evaluate(
        """
        async () => {
          const response = await fetch('/cart/clear.js', {
            method: 'POST',
            credentials: 'include',
            headers: { 'X-Requested-With': 'XMLHttpRequest' }
          });
          const text = await response.text();
          return { ok: response.ok, status: response.status, text };
        }
        """
    )


def read_cart(page: Any) -> Dict[str, Any]:
    result = page.evaluate(
        """
        async () => {
          const response = await fetch('/cart.js', {
            method: 'GET',
            credentials: 'include',
            headers: { 'X-Requested-With': 'XMLHttpRequest' }
          });
          const text = await response.text();
          return { ok: response.ok, status: response.status, text };
        }
        """
    )
    if not isinstance(result, dict) or not result.get("ok"):
        raise ValueError(f"读取购物车失败: HTTP {result.get('status') if isinstance(result, dict) else '-'}")
    try:
        cart = json.loads(str(result.get("text") or "{}"))
    except json.JSONDecodeError as exc:
        raise ValueError("读取购物车失败: /cart.js 返回不是 JSON") from exc
    if not isinstance(cart, dict):
        raise ValueError("读取购物车失败: /cart.js 返回格式异常")
    return cart


def cart_item_title(item: Dict[str, Any]) -> str:
    return str(
        item.get("product_title")
        or item.get("title")
        or item.get("name")
        or item.get("handle")
        or ""
    )


def validate_cart_for_product(cart: Dict[str, Any], product: ProductConfig, *, allow_additional_items: bool = False) -> Dict[str, Any]:
    items = cart.get("items")
    if not isinstance(items, list) or not items:
        raise ValueError("购物车为空，停止支付")
    if not allow_additional_items and len(items) != 1:
        raise ValueError(f"购物车商品数量不是 1（当前 {len(items)}），停止支付")

    matching_items = [item for item in items if int(item.get("variant_id") or 0) == product.variant_id]
    if not matching_items:
        raise ValueError(f"购物车缺少目标 variant_id={product.variant_id}，停止支付")
    item = matching_items[0]
    actual_product_id = int(item.get("product_id") or 0)
    if actual_product_id != product.product_id:
        raise ValueError(f"购物车 product_id 不匹配: expected={product.product_id}, actual={actual_product_id}")
    actual_quantity = int(item.get("quantity") or 0)
    if actual_quantity != product.quantity:
        raise ValueError(f"购物车数量不匹配: expected={product.quantity}, actual={actual_quantity}")
    if product.expected_title_contains and not title_matches(cart_item_title(item), product.expected_title_contains):
        raise ValueError(
            "购物车商品标题不匹配: "
            f"expected contains={product.expected_title_contains}, actual={cart_item_title(item)}"
        )
    return item


def validate_checkout_page_for_product(page: Any, product: ProductConfig) -> None:
    if not product.expected_title_contains:
        return
    body = compact_text(page.text_content("body") or "", 4000)
    if not title_matches(body, product.expected_title_contains):
        raise ValueError(f"checkout 页面未找到目标商品标题: {product.expected_title_contains}")


def detect_verification_challenge(url: str, text: str) -> bool:
    lowered = f"{url or ''}\n{text or ''}".lower()
    keywords = [
        "captcha",
        "hcaptcha",
        "recaptcha",
        "verify you are human",
        "robot",
        "bot protection",
        "人間であること",
        "本人確認",
    ]
    return any(keyword in lowered for keyword in keywords)


def click_pay_button(page: Any) -> None:
    selectors = [
        "#checkout-pay-button",
        "button[type='submit'][name='button']",
        "button:has-text('今すぐ支払う')",
        "button:has-text('支払う')",
        "button:has-text('Pay now')",
    ]
    last_error: Optional[Exception] = None
    for selector in selectors:
        locator = page.locator(selector)
        try:
            if locator.count():
                locator.first.click(timeout=30000)
                return
        except Exception as exc:
            last_error = exc
    if last_error:
        raise RuntimeError(f"点击支付按钮失败: {last_error}") from last_error
    raise RuntimeError("未找到支付按钮")


def wait_best_effort(page: Any, timeout_ms: int = 15000) -> None:
    try:
        page.wait_for_load_state("networkidle", timeout=timeout_ms)
    except Exception:
        try:
            page.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
        except Exception:
            pass


def failure(row: AccountCookie, status: str, message: str, final_url: str = "", order_no: str = "") -> Dict[str, Any]:
    return {
        "row": row.row,
        "email": row.email,
        "success": False,
        "status": status,
        "message": compact_text(message, 800),
        "orderNo": order_no,
        "finalUrl": final_url,
    }


def success(row: AccountCookie, status: str, message: str, final_url: str = "", order_no: str = "") -> Dict[str, Any]:
    return {
        "row": row.row,
        "email": row.email,
        "success": True,
        "status": status,
        "message": compact_text(message, 800),
        "orderNo": order_no,
        "finalUrl": final_url,
    }


def run_account(row: AccountCookie, product: ProductConfig, config: Dict[str, Any]) -> Dict[str, Any]:
    profile = config["profile"]
    headless = bool(config.get("headless", True))
    submit_payment = bool(config.get("submitPayment", True))
    retry_count = max(int(config.get("addToCartRetries", 3) or 0), 0)
    total_add_attempts = retry_count + 1
    label = f"row={row.row} email={row.email or '-'}"
    playwright_manager = None
    browser = None

    try:
        playwright_manager = launch_browser_context()
        playwright = playwright_manager.__enter__()
        browser = create_browser(playwright, headless=headless)
        context = browser.new_context(user_agent=DEFAULT_USER_AGENT, locale="ja-JP")
        cookie_count = add_cookie_header_to_context(context, row.cookie_header)
        if cookie_count <= 0:
            return failure(row, "failed", "Cookie 为空或无法解析")
        page = context.new_page()
        log(f"[{label}] cookieCount={cookie_count} open product")
        page.goto(product.product_url, wait_until="domcontentloaded", timeout=120000)
        wait_best_effort(page)
        if bool(config.get("clearCartBeforeAdd", True)):
            clear_result = clear_cart(page)
            if not isinstance(clear_result, dict) or not clear_result.get("ok"):
                return failure(
                    row,
                    "cart_clear_failed",
                    f"清空购物车失败: HTTP {clear_result.get('status') if isinstance(clear_result, dict) else '-'}",
                    page.url,
                )
            log(f"[{label}] cart cleared")

        last_add_message = ""
        for attempt in range(1, total_add_attempts + 1):
            if attempt > 1:
                time.sleep(random.uniform(0.5, 1.8))
                try:
                    page.goto(product.product_url, wait_until="domcontentloaded", timeout=120000)
                    wait_best_effort(page, timeout_ms=8000)
                except Exception:
                    pass
            try:
                add_result = add_to_cart(page, product)
                if isinstance(add_result, dict) and add_result.get("ok"):
                    log(f"[{label}] add-to-cart ok attempt={attempt}/{total_add_attempts}")
                    break
                last_add_message = f"HTTP {add_result.get('status') if isinstance(add_result, dict) else '-'} {compact_text((add_result or {}).get('text') if isinstance(add_result, dict) else add_result, 300)}"
                log(f"[{label}] add-to-cart failed attempt={attempt}/{total_add_attempts}: {last_add_message}")
            except Exception as exc:
                last_add_message = str(exc)
                log(f"[{label}] add-to-cart exception attempt={attempt}/{total_add_attempts}: {compact_text(last_add_message, 300)}")
        else:
            return failure(row, "add_to_cart_failed", f"Jump Shop 加购失败: {last_add_message}", page.url)

        try:
            cart = read_cart(page)
            cart_item = validate_cart_for_product(
                cart,
                product,
                allow_additional_items=bool(config.get("allowAdditionalCartItems", False)),
            )
            log(
                f"[{label}] cart verified variantId={cart_item.get('variant_id')} "
                f"productId={cart_item.get('product_id')} quantity={cart_item.get('quantity')} "
                f"title={compact_text(cart_item_title(cart_item), 120)}"
            )
        except Exception as exc:
            return failure(row, "cart_mismatch", f"支付前购物车校验失败: {exc}", page.url)

        log(f"[{label}] open checkout")
        page.goto(JUMP_SHOP_CHECKOUT_URL, wait_until="domcontentloaded", timeout=120000)
        wait_best_effort(page)
        page.wait_for_timeout(2500)
        if "/account/login" in page.url:
            return failure(row, "login_required", "Jump Shop checkout 跳回登录页，Cookie 登录态失效", page.url)

        current_text = compact_text(page.text_content("body") or "", 2000)
        if detect_verification_challenge(page.url, current_text):
            return failure(row, "verification_required", "Jump Shop checkout 命中人机验证", page.url)
        try:
            validate_checkout_page_for_product(page, product)
        except Exception as exc:
            return failure(row, "checkout_mismatch", f"支付前 checkout 商品校验失败: {exc}", page.url)

        if row.email:
            maybe_fill(page, ["input[type='email']", "input[name='email']", "#email"], row.email)
        fill_shipping_address(page, profile)
        fill_billing_address(page, profile)
        fill_card_fields(page, profile)
        log(f"[{label}] checkout fields filled")

        if not submit_payment:
            return success(row, "ready_to_pay", "dry-run: 已加购并填完 checkout，未点击支付", page.url)

        click_pay_button(page)
        try:
            page.wait_for_load_state("domcontentloaded", timeout=120000)
        except Exception:
            pass
        page.wait_for_timeout(6000)
        final_text = compact_text(page.text_content("body") or "", 4000)
        final_url = page.url
        if detect_jump_shop_3ds(final_url, final_text):
            return failure(row, "3ds_required", "3DS_REQUIRED: Jump Shop 命中外部验证页", final_url)
        if detect_verification_challenge(final_url, final_text):
            return failure(row, "verification_required", "Jump Shop 支付后命中人机验证", final_url)
        if detect_jump_shop_success(final_url, final_text):
            order_no = extract_order_number(final_text, final_url)
            return success(row, "paid", "Jump Shop 下单成功", final_url, order_no)
        return failure(row, "payment_failed", f"Jump Shop 下单失败: {final_text[:300] or final_url}", final_url)
    except Exception as exc:
        return failure(row, "exception", f"Jump Shop 执行异常: {exc}", "")
    finally:
        if browser is not None:
            try:
                browser.close()
            except Exception:
                pass
        if playwright_manager is not None:
            try:
                playwright_manager.__exit__(None, None, None)
            except Exception:
                pass


def write_results(path: Path, results: Iterable[Dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=RESULT_FIELDS)
        writer.writeheader()
        for result in sorted(results, key=lambda item: int(item.get("row") or 0)):
            writer.writerow({field: result.get(field, "") for field in RESULT_FIELDS})


def run(config_path: Path) -> int:
    config_path = config_path.resolve()
    config = load_json(config_path)
    base_dir = config_path.parent
    excel_path = resolve_path(base_dir, config.get("excelPath"))
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel 文件不存在: {excel_path}")

    if not isinstance(config.get("profile"), dict):
        raise ValueError("config 缺少 profile")
    validate_profile(config["profile"])
    product = resolve_product(config)
    accounts = load_accounts_from_excel(excel_path, str(config.get("sheetName") or ""))
    if not accounts:
        raise ValueError("Excel 中没有可用 cookie 行")

    concurrency = max(int(config.get("concurrency") or 1), 1)
    concurrency = min(concurrency, len(accounts))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    result_path = base_dir / f"quick_checkout_results_{timestamp}.csv"

    log(
        f"[start] accounts={len(accounts)} concurrency={concurrency} quantity={product.quantity} "
        f"submitPayment={bool(config.get('submitPayment', True))} results={result_path}"
    )

    results: List[Dict[str, Any]] = []
    with ThreadPoolExecutor(max_workers=concurrency) as executor:
        futures = {executor.submit(run_account, account, product, config): account for account in accounts}
        for future in as_completed(futures):
            account = futures[future]
            try:
                result = future.result()
            except Exception as exc:
                result = failure(account, "exception", f"线程异常: {exc}")
            results.append(result)
            log(
                f"[result] row={result.get('row')} email={result.get('email') or '-'} "
                f"success={result.get('success')} status={result.get('status')} message={result.get('message')}"
            )
            write_results(result_path, results)

    success_count = sum(1 for item in results if item.get("success"))
    log(f"[completed] success={success_count} failed={len(results) - success_count} results={result_path}")
    return 0 if success_count > 0 else 2


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Jump Shop Excel Cookie quick checkout runner")
    parser.add_argument("--config", required=True, help="Path to quick_checkout.config.json")
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    try:
        return run(Path(args.config))
    except KeyboardInterrupt:
        log("[aborted] interrupted by user")
        return 130
    except Exception as exc:
        log(f"[fatal] {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
